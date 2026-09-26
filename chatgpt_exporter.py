import json
import re
from io import BytesIO
from typing import Any

import requests
from docx import Document
from docx.shared import Cm, Pt


SHARE_PAGE_BASE = "https://chatgpt.com/share/"


class ChatExportError(RuntimeError):
    """Понятная пользователю ошибка при получении или обработке shared-чата."""


def get_share_id(value: str) -> str:
    """Извлекает share ID из ссылки ChatGPT или принимает сам ID."""
    value = value.strip()
    if not value:
        raise ChatExportError("Вставьте shared-ссылку ChatGPT.")

    match = re.search(
        r"(?:https?://)?(?:www\.)?(?:chatgpt\.com|chat\.openai\.com)/share/([^/?#\s]+)",
        value,
        flags=re.IGNORECASE,
    )
    share_id = match.group(1) if match else value

    if not re.fullmatch(r"[A-Za-z0-9_-]+", share_id):
        raise ChatExportError("Не удалось распознать shared-ссылку ChatGPT.")

    return share_id


def get_message_text(message: dict[str, Any]) -> str:
    content = message.get("content") or {}
    result: list[str] = []

    if isinstance(content, str):
        return content.strip()

    if not isinstance(content, dict):
        return ""

    content_type = content.get("content_type")
    if content_type == "code" and isinstance(content.get("text"), str):
        language = content.get("language") or ""
        return f"```{language}\n{content['text']}\n```".strip()

    parts = content.get("parts", [])
    if isinstance(parts, list):
        for part in parts:
            if isinstance(part, str):
                if part.strip():
                    result.append(part)
            elif isinstance(part, dict):
                text = part.get("text")
                if isinstance(text, str) and text.strip():
                    result.append(text)
                    continue

                part_type = part.get("content_type") or part.get("type")
                if part_type in ("image_asset_pointer", "image"):
                    result.append("[изображение]")
                elif part_type in ("audio_asset_pointer", "audio_transcription"):
                    result.append("[аудио]")
                elif part_type == "file":
                    result.append("[файл]")

    direct_text = content.get("text")
    if not result and isinstance(direct_text, str) and direct_text.strip():
        result.append(direct_text)

    return "\n".join(result).strip()


def _message_is_exportable(message: dict[str, Any]) -> bool:
    author = message.get("author") or {}
    role = author.get("role")
    if role not in ("user", "assistant"):
        return False

    metadata = message.get("metadata") or {}
    if metadata.get("is_visually_hidden_from_conversation"):
        return False

    recipient = message.get("recipient")
    if role == "assistant" and recipient not in (None, "all"):
        return False

    return True


def get_messages_from_linear(data: dict[str, Any]) -> list[tuple[str, str]]:
    messages: list[tuple[str, str]] = []

    for item in data.get("linear_conversation", []):
        if not isinstance(item, dict):
            continue

        message = item.get("message")
        if not message and "author" in item:
            message = item

        if not isinstance(message, dict) or not _message_is_exportable(message):
            continue

        role = (message.get("author") or {}).get("role")
        text = get_message_text(message)
        if text:
            messages.append((role, text))

    return messages


def get_messages_from_mapping(data: dict[str, Any]) -> list[tuple[str, str]]:
    mapping = data.get("mapping") or {}
    if not isinstance(mapping, dict) or not mapping:
        return []

    current = data.get("current_node")
    nodes: list[dict[str, Any]] = []

    if current and current in mapping:
        while current:
            node = mapping.get(current)
            if not isinstance(node, dict):
                break
            nodes.append(node)
            current = node.get("parent")
        nodes.reverse()
    else:
        roots = [
            node_id
            for node_id, node in mapping.items()
            if isinstance(node, dict) and not node.get("parent")
        ]

        if roots:
            seen: set[str] = set()

            def walk(node_id: str) -> None:
                if node_id in seen:
                    return
                node = mapping.get(node_id)
                if not isinstance(node, dict):
                    return
                seen.add(node_id)
                nodes.append(node)
                for child_id in node.get("children") or []:
                    if isinstance(child_id, str):
                        walk(child_id)

            for root_id in roots:
                walk(root_id)
        else:
            nodes = [node for node in mapping.values() if isinstance(node, dict)]

            def message_time(node: dict[str, Any]) -> float:
                message = node.get("message") or {}
                return message.get("create_time") or 0

            nodes.sort(key=message_time)

    messages: list[tuple[str, str]] = []
    for node in nodes:
        message = node.get("message")
        if not isinstance(message, dict) or not _message_is_exportable(message):
            continue

        role = (message.get("author") or {}).get("role")
        text = get_message_text(message)
        if text:
            messages.append((role, text))

    return messages


# -----------------------------------------------------------------------------
# Public share page decoder
# -----------------------------------------------------------------------------


def _read_js_string(source: str, quote_index: int) -> tuple[str, int]:
    """Читает тело JS-строки, сохраняя escape-последовательности."""
    if quote_index >= len(source) or source[quote_index] != '"':
        raise ValueError("Ожидалась JS-строка.")

    i = quote_index + 1
    out: list[str] = []

    while i < len(source):
        char = source[i]
        if char == "\\":
            if i + 1 < len(source):
                out.append(source[i : i + 2])
                i += 2
                continue
        if char == '"':
            return "".join(out), i + 1
        out.append(char)
        i += 1

    return "".join(out), i


def _extract_turbo_stream(page_html: str) -> str:
    """Собирает React Router hydration payload из streamController.enqueue(...)."""
    chunks: list[str] = []

    for match in re.finditer(r"streamController\.enqueue\(", page_html):
        pos = match.end()
        while pos < len(page_html) and page_html[pos] in " \r\n\t":
            pos += 1

        if pos >= len(page_html) or page_html[pos] != '"':
            continue

        raw, _ = _read_js_string(page_html, pos)
        try:
            chunks.append(json.loads('"' + raw + '"'))
        except json.JSONDecodeError:
            continue

    return "".join(chunks)


class _TurboDecoder:
    def __init__(self, flat: list[Any], promises: dict[int, Any] | None = None):
        self.flat = flat
        self.promises = promises or {}
        self.memo: dict[int, Any] = {}
        self.in_progress: set[int] = set()

    def resolve_edge(self, edge: Any) -> Any:
        if isinstance(edge, bool):
            return edge
        if isinstance(edge, int):
            if edge < 0:
                return None
            if edge in self.promises:
                return self.promises[edge]
            if 0 <= edge < len(self.flat):
                return self.resolve_index(edge)
            return None
        return edge

    def resolve_index(self, index: int) -> Any:
        if index in self.memo:
            return self.memo[index]
        if index in self.in_progress:
            return None

        self.in_progress.add(index)
        node = self.flat[index]

        if isinstance(node, dict):
            value: dict[str, Any] = {}
            self.memo[index] = value

            for raw_key, raw_value in node.items():
                key: Any = raw_key
                if isinstance(raw_key, str) and raw_key.startswith("_"):
                    suffix = raw_key[1:]
                    if suffix.lstrip("-").isdigit():
                        key = self.resolve_edge(int(suffix))

                if not isinstance(key, str):
                    key = str(key)

                value[key] = self.resolve_edge(raw_value)

        elif isinstance(node, list):
            value = []
            self.memo[index] = value
            value.extend(self.resolve_edge(item) for item in node)
        else:
            value = node

        self.memo[index] = value
        self.in_progress.discard(index)
        return value


def _decode_promise_lines(lines: list[str]) -> dict[int, Any]:
    promises: dict[int, Any] = {}

    for line in lines:
        match = re.match(r"^P(\d+):(.*)$", line)
        if not match:
            continue

        promise_index = int(match.group(1))
        body = match.group(2)

        try:
            sub_flat = json.loads(body)
        except json.JSONDecodeError:
            continue

        if isinstance(sub_flat, list) and sub_flat:
            try:
                promises[promise_index] = _TurboDecoder(sub_flat).resolve_index(0)
            except Exception:
                promises[promise_index] = sub_flat
        else:
            promises[promise_index] = sub_flat

    return promises


def _decode_turbo_stream(stream: str) -> Any:
    if not stream or not stream.strip():
        raise ValueError("Пустой hydration payload.")

    lines = stream.split("\n")
    flat = json.loads(lines[0])
    if not isinstance(flat, list):
        raise ValueError("Неожиданный формат hydration payload.")

    promises = _decode_promise_lines(lines[1:])
    return _TurboDecoder(flat, promises).resolve_index(0)


def _find_conversation_payload(root: Any) -> dict[str, Any] | None:
    """Ищет объект разговора по форме, а не по жёсткому пути."""
    if not isinstance(root, (dict, list)):
        return None

    stack: list[tuple[Any, int]] = [(root, 0)]
    seen: set[int] = set()

    while stack:
        node, depth = stack.pop()
        if depth > 60:
            continue

        if isinstance(node, dict):
            node_id = id(node)
            if node_id in seen:
                continue
            seen.add(node_id)

            mapping = node.get("mapping")
            linear = node.get("linear_conversation")
            if isinstance(mapping, dict) or isinstance(linear, list):
                return node

            for value in node.values():
                if isinstance(value, (dict, list)):
                    stack.append((value, depth + 1))

        elif isinstance(node, list):
            for value in reversed(node):
                if isinstance(value, (dict, list)):
                    stack.append((value, depth + 1))

    return None


def _extract_json_script_candidates(page_html: str) -> list[Any]:
    """Fallback для старых/альтернативных страниц с JSON внутри <script>."""
    candidates: list[Any] = []

    patterns = [
        r'<script[^>]+id=["\']__NEXT_DATA__["\'][^>]*>(.*?)</script>',
        r'<script[^>]+type=["\']application/json["\'][^>]*>(.*?)</script>',
    ]

    for pattern in patterns:
        for match in re.finditer(pattern, page_html, flags=re.IGNORECASE | re.DOTALL):
            raw = match.group(1).strip()
            if not raw:
                continue
            try:
                candidates.append(json.loads(raw))
            except json.JSONDecodeError:
                continue

    return candidates


def _parse_conversation_from_html(page_html: str) -> dict[str, Any]:
    # Современный ChatGPT share page: React Router / turbo-stream.
    stream = _extract_turbo_stream(page_html)
    if stream:
        try:
            root = _decode_turbo_stream(stream)
            conversation = _find_conversation_payload(root)
            if conversation is not None:
                return conversation
        except (ValueError, json.JSONDecodeError, TypeError, RecursionError):
            pass

    # Fallback: старые страницы или иной JSON bootstrap.
    for candidate in _extract_json_script_candidates(page_html):
        conversation = _find_conversation_payload(candidate)
        if conversation is not None:
            return conversation

    raise ChatExportError(
        "Публичная страница чата загрузилась, но структуру разговора распознать не удалось. "
        "Возможно, ChatGPT снова изменил формат share-страницы."
    )


def fetch_shared_chat(shared_link: str) -> tuple[str, list[tuple[str, str]]]:
    """
    Загружает публичную share-страницу ChatGPT и извлекает название и сообщения.

    Важно: это не официальный API. Формат публичной страницы может измениться.
    """
    share_id = get_share_id(shared_link)
    page_url = SHARE_PAGE_BASE + share_id

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/153.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "ru,en;q=0.9",
    }

    try:
        response = requests.get(
            page_url,
            headers=headers,
            timeout=60,
            allow_redirects=True,
        )
    except requests.Timeout as exc:
        raise ChatExportError("ChatGPT слишком долго не отвечает. Попробуйте ещё раз.") from exc
    except requests.RequestException as exc:
        raise ChatExportError("Не удалось подключиться к публичной странице ChatGPT.") from exc

    if response.status_code == 404:
        raise ChatExportError(
            "Shared-чат не найден. Проверьте ссылку и убедитесь, что общий доступ к чату включён."
        )

    if response.status_code == 403:
        raise ChatExportError(
            "ChatGPT не разрешил серверу открыть публичную share-страницу (403). "
            "Этот способ нельзя надёжно использовать с текущего хостинга."
        )

    if response.status_code != 200:
        raise ChatExportError(f"ChatGPT вернул ошибку HTTP {response.status_code}.")

    conversation = _parse_conversation_from_html(response.text)

    title = conversation.get("title") or "ChatGPT conversation"

    messages: list[tuple[str, str]] = []
    if conversation.get("linear_conversation"):
        messages = get_messages_from_linear(conversation)
    if not messages and conversation.get("mapping"):
        messages = get_messages_from_mapping(conversation)

    if not messages:
        raise ChatExportError(
            "Чат загрузился, но текстовые сообщения не найдены. "
            "Возможно, формат share-страницы изменился."
        )

    return str(title), messages


def safe_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]', "_", name).strip(" .")
    return name[:150] or "ChatGPT conversation"


def build_docx(title: str, messages: list[tuple[str, str]]) -> bytes:
    """Собирает Word-документ в памяти и возвращает его как bytes."""
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    title_paragraph = doc.add_heading(title, level=0)
    title_paragraph.paragraph_format.space_after = Pt(18)

    for index, (role, text) in enumerate(messages):
        speaker = "Вы" if role == "user" else "ChatGPT"

        speaker_paragraph = doc.add_paragraph()
        speaker_paragraph.paragraph_format.space_before = Pt(6)
        speaker_paragraph.paragraph_format.space_after = Pt(4)
        speaker_run = speaker_paragraph.add_run(speaker)
        speaker_run.bold = True
        speaker_run.font.size = Pt(12)

        body = doc.add_paragraph(text)
        body.paragraph_format.space_after = Pt(10)
        body.paragraph_format.line_spacing = 1.08

        if index != len(messages) - 1:
            separator = doc.add_paragraph()
            separator.paragraph_format.space_after = Pt(3)
            separator.add_run("─" * 48)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_shared_chat(shared_link: str) -> tuple[str, str, int, bytes]:
    """Полный сценарий: публичная share-ссылка -> чат -> DOCX."""
    title, messages = fetch_shared_chat(shared_link)
    filename = safe_filename(title) + ".docx"
    docx_bytes = build_docx(title, messages)
    return title, filename, len(messages), docx_bytes
